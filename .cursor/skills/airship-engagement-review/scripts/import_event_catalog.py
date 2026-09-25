#!/usr/bin/env python3
"""Import the Airship "Data Collection by Vertical" workbook into event_catalog.json
(+ event_catalog.md).

Usage:
    python scripts/import_event_catalog.py ["Data Collection by Vertical.xlsx"]

If no path is given, defaults to ~/Downloads/Data Collection by Vertical.xlsx.

The workbook ships one "event categories" sheet (the canonical category list) plus one
sheet per vertical (Travel, Media, Sport & Gaming, Retail, FinTech, EV Charging & Fuel,
Casino & Gaming, Recipes & Cooking) and an "all verticals" sheet of common standards.
Each vertical sheet has two sections:
  - Attributes    : Name | ID | (blank) | Description | Type | Source
  - Custom Events : Category | custom_name | Properties | Description | Types | Source

Column layout (1-based): A blank · B Category/Name · C custom_name/ID · D Properties ·
E Description · F Types · G Source. An event group starts on the row carrying B+C; each
following property row carries only D (+F type) until the next B+C row.

This catalog powers the custom-event contextualisation in the engagement review:
classify a client's events against the vertical's recommended events, flag conversion
KPIs (business/usage), know which events *should* carry a monetary value/currency, and
surface data-collection opportunities. Machine-readable copy is `event_catalog.json`
(scripts read that); `event_catalog.md` is the human reference.
"""
import datetime
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)

DEFAULT_XLSX = os.path.expanduser("~/Downloads/Data Collection by Vertical.xlsx")

# Map an engagement-review benchmark vertical (benchmarks.json slug / client industry)
# to the best-fit vertical sheet in this workbook. Anything unmapped falls back to
# "all_verticals" (handled by the consumer, e.g. event_analysis.resolve_vertical).
VERTICAL_MAP = {
    "retail": "retail",
    "media": "media",
    "entertainment": "media",
    "finance_insurance": "fintech",
    "gambling_gaming": "casino_gaming",
    "sports_recreation": "sport_gaming",
    "travel_transportation": "travel",
    "food_drink": "recipes_cooking",
    # utility_productivity / business / government / education / social / health /
    # charities have no dedicated sheet -> all_verticals fallback.
}

# Per-category KPI semantics used across the review. kpi_kind: "business" (revenue /
# transaction / acquisition), "loyalty" (points/tier), "usage" (consumption / feature
# use / engagement). is_conversion flags a category whose occurrence is a meaningful
# conversion or consumption KPI (worth attribution analysis).
CATEGORY_KPI = {
    "purchase": ("business", True),
    "transaction success": ("business", True),
    "add to cart": ("business", True),
    "acquisition": ("business", True),
    "activation": ("business", True),
    "offer": ("business", True),
    "preorder": ("business", True),
    "reservation": ("business", True),
    "loyalty": ("loyalty", True),
    "consume content/product": ("usage", True),
    "feature usage": ("usage", False),
    "engagement": ("usage", False),
    "in-session experience": ("usage", False),
    "save content/product": ("usage", False),
    "share content/product": ("usage", False),
    "star content/product": ("usage", False),
    "browse content/product": ("usage", False),
    "view": ("usage", False),
    "add to wishlist": ("usage", False),
    "update wishlist": ("usage", False),
    "remove from cart": ("usage", False),
    "remove from wishlist": ("usage", False),
    "update cart": ("business", False),
    "account registration": ("usage", False),
    "login": ("usage", False),
    "logout": ("usage", False),
    "update profile": ("usage", False),
    "location/geofence": ("usage", False),
    "transaction failure": ("usage", False),
    "cancellation": ("usage", False),
    "confirmation": ("usage", False),
    "delivery": ("usage", False),
    "expiration": ("usage", False),
    "click": ("usage", False),
    "ad display": ("usage", False),
}

# A property is treated as a monetary amount when its (lowercased) name matches this.
MONEY_PROP_RE = re.compile(
    r"(?:^|_)(amount|price|value)$"
    r"|total_amount|total_win_loss|final_win_loss_amount|win_loss"
    r"|bet_amount|transaction_value|check_amount|credit_limit"
    r"|credits_value|price_usd|original_price|new_reduced_price",
    re.I,
)
# Preference order when several monetary properties exist on one event.
MONEY_PRIORITY = [
    "total_amount", "amount", "value", "transaction_value", "bet_amount",
    "check_amount", "final_win_loss_amount", "total_win_loss",
]


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", str(s).strip().lower()).strip("_")


def _cell(row, idx):
    v = row[idx] if len(row) > idx else None
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def is_money_prop(name):
    return bool(name) and bool(MONEY_PROP_RE.search(name.strip().lower()))


def pick_value_property(properties):
    """Return the best monetary property name for an event (or None)."""
    names = [p["name"] for p in properties]
    lowered = {n.lower(): n for n in names}
    for pref in MONEY_PRIORITY:
        if pref in lowered:
            return lowered[pref]
    for n in names:
        if is_money_prop(n):
            return n
    return None


def category_flags(category):
    return CATEGORY_KPI.get((category or "").strip().lower(), ("usage", False))


def parse_categories(wb):
    if "event categories" not in wb.sheetnames:
        return []
    ws = wb["event categories"]
    out = []
    for row in ws.iter_rows(values_only=True):
        b = _cell(row, 1)
        if not b or b.lower() == "event categories":
            continue
        out.append(b)
    return out


def parse_vertical(ws):
    """Parse one vertical sheet into {attributes: [...], events: [...]}."""
    attributes = []
    events = []
    mode = None
    current = None
    for row in ws.iter_rows(values_only=True):
        b, c, d, e, f = (_cell(row, i) for i in (1, 2, 3, 4, 5))
        if b == "Attributes":
            mode = "attr"
            continue
        if b == "Custom Events":
            mode = "cust"
            continue
        if b in ("Name", "Category"):  # section header row
            continue
        if mode == "attr":
            if b and c:
                attributes.append({"name": b, "id": c, "type": (f or "").lower() or None,
                                   "description": e})
            continue
        if mode == "cust":
            if b and c:  # new event group
                current = {
                    "custom_name": c,
                    "category": b,
                    "description": e,
                    "properties": [],
                }
                if d:
                    current["properties"].append({"name": d, "type": (f or "").lower() or None})
                events.append(current)
            elif d and current is not None:  # continuation property row
                current["properties"].append({"name": d, "type": (f or "").lower() or None})
    # enrich derived flags
    for ev in events:
        kind, is_conv = category_flags(ev["category"])
        ev["kpi_kind"] = kind
        ev["is_conversion"] = is_conv
        ev["value_property"] = pick_value_property(ev["properties"])
        ev["has_currency_property"] = any(
            p["name"].strip().lower() == "currency" for p in ev["properties"]
        )
    return {"attributes": attributes, "events": events}


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx):
        sys.exit(f"Workbook not found: {xlsx}")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    categories = parse_categories(wb)
    verticals = {}
    for name in wb.sheetnames:
        if name == "event categories":
            continue
        key = "all_verticals" if slug(name) in ("all_verticals", "all_vertical") else slug(name)
        parsed = parse_vertical(wb[name])
        parsed["label"] = name
        parsed["key"] = key
        verticals[key] = parsed

    out = {
        "meta": {
            "source": "Airship — Data Collection by Vertical (best-practice event & attribute catalog)",
            "file": os.path.basename(xlsx),
            "imported": datetime.date.today().isoformat(),
            "notes": (
                "Best-practice custom events & attributes per vertical. 'Category' is the "
                "use-case grouping used in the workbook (Purchase, Loyalty, Engagement, "
                "In-session Experience, ...), distinct from the canonical 'categories' list. "
                "value_property = the event's monetary property (amount/value/price/...) when "
                "one exists; has_currency_property flags an explicit 'currency' property. "
                "kpi_kind/is_conversion are derived from the category (see import script). "
                "This catalog is guidance, not the client's data."
            ),
        },
        "categories": categories,
        "vertical_map": VERTICAL_MAP,
        "verticals": dict(sorted(verticals.items())),
    }

    json_path = os.path.join(SKILL_DIR, "event_catalog.json")
    with open(json_path, "w") as fh:
        json.dump(out, fh, indent=2, ensure_ascii=False)
    n_events = sum(len(v["events"]) for v in verticals.values())
    print(f"Wrote {json_path}: {len(verticals)} verticals, {n_events} events, "
          f"{len(categories)} canonical categories")

    write_markdown(out)


def write_markdown(out):
    m = out["meta"]
    lines = [
        "# Custom-event catalog by vertical (engagement review)",
        "",
        f"Source: **{m['source']}** · file `{m['file']}` · imported {m['imported']}.",
        "",
        "Machine-readable copy: `event_catalog.json` (scripts read that). Regenerate with",
        "`python scripts/import_event_catalog.py \"Data Collection by Vertical.xlsx\"`.",
        "",
        "## How the skill uses this file",
        "1. Resolve the client's vertical (from benchmarks industry) to a book vertical via",
        "   `vertical_map` (fallback `all_verticals`).",
        "2. `event_analysis.classify_event()` matches each client custom event to a catalog",
        "   event/category with a confidence level (name tokens + aliases).",
        "3. Conversion KPIs are flagged from the matched category (`is_conversion`, `kpi_kind`).",
        "4. `value_property` / `has_currency_property` tell whether an event *should* carry a",
        "   monetary amount — used to spot 'sent but without conversion value' opportunities.",
        "5. Catalog events the client does NOT send become 'send more to Airship' opportunities.",
        "",
        f"> Notes: {m['notes']}",
        "",
        "## Vertical mapping (benchmark industry -> book vertical)",
        "| Benchmark vertical | Book vertical |",
        "|---|---|",
    ]
    for k, v in out["vertical_map"].items():
        lines.append(f"| `{k}` | `{v}` |")
    lines += ["| _(any other)_ | `all_verticals` |", "",
              "## Canonical event categories", ""]
    lines.append(", ".join(f"`{c}`" for c in out["categories"]))
    lines += ["", "## Recommended events by vertical", ""]
    for key, v in out["verticals"].items():
        lines.append(f"### {v['label']}  (`{key}`)")
        lines.append("")
        lines.append("| Event | Category | KPI kind | Conversion | Value prop | Currency |")
        lines.append("|---|---|---|---|---|---|")
        for ev in v["events"]:
            lines.append(
                f"| `{ev['custom_name']}` | {ev['category']} | {ev['kpi_kind']} | "
                f"{'yes' if ev['is_conversion'] else '—'} | "
                f"{('`'+ev['value_property']+'`') if ev['value_property'] else '—'} | "
                f"{'yes' if ev['has_currency_property'] else '—'} |"
            )
        lines.append("")
    md_path = os.path.join(SKILL_DIR, "event_catalog.md")
    with open(md_path, "w") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"Wrote {md_path}")


if __name__ == "__main__":
    main()
